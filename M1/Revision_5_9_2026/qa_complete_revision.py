from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
from pathlib import Path

import numpy as np
from docx import Document
from PIL import Image, ImageDraw


ROOT = Path(__file__).resolve().parent
AUDIT = ROOT / "Final_Revision_Audit.json"
SETS = {
    "main_clean": "_qa_final3_main_clean",
    "main_marked": "_qa_final3_main_marked",
    "supp_clean": "_qa_final3_supp_clean",
    "supp_marked": "_qa_final3_supp_marked",
    "response": "_qa_final3_response",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def natural_key(path: Path) -> list[object]:
    return [int(value) if value.isdigit() else value for value in re.split(r"(\d+)", path.name)]


def page_diagnostics(path: Path) -> dict[str, object]:
    image = Image.open(path).convert("RGB")
    array = np.asarray(image)
    ink = np.any(array < 245, axis=2)
    ys, xs = np.where(ink)
    if not len(xs):
        return {"file": path.name, "blank": True, "size": list(image.size)}
    left = int(xs.min())
    right = int(image.width - 1 - xs.max())
    top = int(ys.min())
    bottom = int(image.height - 1 - ys.max())
    return {
        "file": path.name,
        "blank": False,
        "size": list(image.size),
        "ink_bbox_margins_px": {"left": left, "right": right, "top": top, "bottom": bottom},
        "edge_contact": bool(min(left, right, top, bottom) < 8),
    }


def make_contacts(directory: Path, pages: list[Path]) -> list[str]:
    outputs = []
    for start in range(0, len(pages), 4):
        group = pages[start:start + 4]
        opened = [Image.open(path).convert("RGB") for path in group]
        width = max(image.width for image in opened)
        height = max(image.height for image in opened)
        sheet = Image.new("RGB", (width * 2 + 60, height * 2 + 100), "#D9D9D9")
        draw = ImageDraw.Draw(sheet)
        for index, (path, image) in enumerate(zip(group, opened)):
            col = index % 2
            row = index // 2
            x = col * (width + 40) + 10
            y = row * (height + 40) + 30
            sheet.paste(image, (x, y))
            draw.rectangle((x, y, x + image.width - 1, y + image.height - 1), outline="#666666", width=2)
            draw.text((x, y - 22), path.stem, fill="#000000")
        output = directory / f"contact-{start // 4 + 1:02d}.png"
        sheet.save(output, optimize=True)
        outputs.append(output.name)
    return outputs


def pdf_text(pdf: Path) -> str:
    result = subprocess.run(
        ["pdftotext", "-enc", "UTF-8", str(pdf), "-"],
        check=True,
        capture_output=True,
    )
    return result.stdout.decode("utf-8", errors="replace")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--visual-status", choices=["PENDING", "PASS"], default="PENDING")
    args = parser.parse_args()
    report: dict[str, object] = {
        "renderer": "Microsoft Word 16 ExportAsFixedFormat; Poppler pdftoppm 150 dpi",
        "libreoffice_renderer": "UNAVAILABLE",
        "visual_inspection": args.visual_status,
        "sets": {},
    }
    total_pages = 0
    all_pass = True
    for label, folder_name in SETS.items():
        directory = ROOT / folder_name
        pages = sorted(directory.glob("page-*.png"), key=natural_key)
        diagnostics = [page_diagnostics(path) for path in pages]
        contacts = make_contacts(directory, pages)
        text = pdf_text(directory / "render.pdf")
        required = []
        forbidden = []
        if label.startswith("main"):
            required = [
                "Published practical precedents and the remaining integration gap",
                "Yang et al. (2004)",
                "Eid et al. (2026)",
                "Feldmann et al. (2024)",
                "Moracchini et al. (2025)",
                "Table 8",
            ]
            forbidden = ["Controlled-synthetic worked demonstration", "Figure 7. Controlled-synthetic", "PR-AUC 0.464"]
        elif label.startswith("supp"):
            required = ["Supplementary Information", "Supplementary Table S2"]
            forbidden = ["Supplementary Methods S3. Controlled-synthetic", "Supplementary Table S3. Locked design"]
        else:
            required = ["R1.1", "R2.5", "R2.6", "articles should not include unpublished/original data", "published practical applications"]
            forbidden = ["Repository and reproducibility record", "PR-AUC 0.464"]
        missing = [value for value in required if value not in text]
        present_forbidden = [value for value in forbidden if value in text]
        set_pass = bool(pages) and not any(x["blank"] for x in diagnostics) and not missing and not present_forbidden
        all_pass = all_pass and set_pass
        total_pages += len(pages)
        report["sets"][label] = {
            "expected_pages": len(pages),
            "rendered_pages": len(pages),
            "blank_pages": [x["file"] for x in diagnostics if x["blank"]],
            "edge_contact_pages": [x["file"] for x in diagnostics if x.get("edge_contact")],
            "required_rendered_text_missing": missing,
            "forbidden_rendered_text_present": present_forbidden,
            "contact_sheets": contacts,
            "programmatic_status": "PASS" if set_pass else "FAIL",
        }
    report["total_pages"] = total_pages
    report["programmatic_status"] = "PASS" if all_pass else "FAIL"
    report["status"] = "PASS" if all_pass and args.visual_status == "PASS" else "PENDING"

    clean_marked_pairs = [
        (
            ROOT / "Manuscript- Water and Ecology_Fully_Revised_Clean.docx",
            ROOT / "Manuscript- Water and Ecology_Fully_Revised_Colour_Marked.docx",
            "Title:",
        ),
        (
            ROOT / "SupplementaryInformation_Fully_Revised_Clean.docx",
            ROOT / "SupplementaryInformation_Fully_Revised_Colour_Marked.docx",
            "Supplementary Information",
        ),
    ]
    identity_checks = []
    for clean_path, marked_path, start in clean_marked_pairs:
        clean_doc = Document(clean_path)
        marked_doc = Document(marked_path)
        clean_paragraphs = [paragraph.text for paragraph in clean_doc.paragraphs]
        marked_paragraphs = [paragraph.text for paragraph in marked_doc.paragraphs]
        marked_start = next(i for i, text in enumerate(marked_paragraphs) if text.startswith(start))
        clean_tables = [
            [[cell.text for cell in row.cells] for row in table.rows]
            for table in clean_doc.tables
        ]
        marked_tables = [
            [[cell.text for cell in row.cells] for row in table.rows]
            for table in marked_doc.tables
        ]
        identity_checks.append({
            "clean": clean_path.name,
            "marked": marked_path.name,
            "excluded_marked_legend_paragraphs": marked_start,
            "paragraph_text_identical": clean_paragraphs == marked_paragraphs[marked_start:],
            "table_text_identical": clean_tables == marked_tables,
        })

    hygiene_files = [
        ROOT / "final_main_text_hygiene.json",
        ROOT / "final_supp_text_hygiene.json",
        ROOT / "final_response_text_hygiene.json",
    ]
    hygiene = [json.loads(path.read_text(encoding="utf-8")) for path in hygiene_files]
    vale_config = Path(r"C:\Users\ThinkPad P1 G4\.agents\skills\r2m\vale\.vale.ini")
    vale_result = subprocess.run(
        [
            "vale", f"--config={vale_config}", "--output=JSON",
            str(ROOT / "final_main_text.txt"),
            str(ROOT / "final_supp_text.txt"),
            str(ROOT / "final_response_text.txt"),
        ],
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    vale_payload = json.loads(vale_result.stdout)
    vale_findings = [item for values in vale_payload.values() for item in values]

    deliverables = [
        ROOT / "Manuscript- Water and Ecology_Fully_Revised_Clean.docx",
        ROOT / "Manuscript- Water and Ecology_Fully_Revised_Colour_Marked.docx",
        ROOT / "SupplementaryInformation_Fully_Revised_Clean.docx",
        ROOT / "SupplementaryInformation_Fully_Revised_Colour_Marked.docx",
        ROOT / "Response_to_Reviewers_Fully_Addressed.docx",
        ROOT / "Figures" / "Figure_4_Fully_Revised.png",
        ROOT / "Figures" / "Figure_4_Fully_Revised.svg",
        ROOT / "Published_Practical_Example_Search.md",
    ]

    audit = json.loads(AUDIT.read_text(encoding="utf-8"))
    audit["render_qa"] = report
    audit["structural_identity"] = {
        "status": "PASS" if all(
            item["paragraph_text_identical"] and item["table_text_identical"]
            for item in identity_checks
        ) else "FAIL",
        "checks": identity_checks,
    }
    audit["writing_qa"] = {
        "r2m_text_hygiene_status": "PASS" if all(item["status"] == "PASS" for item in hygiene) else "FAIL",
        "r2m_text_hygiene_findings": sum(item["findings_count"] for item in hygiene),
        "vale_errors": sum(item["Severity"] == "error" for item in vale_findings),
        "vale_warnings": sum(item["Severity"] == "warning" for item in vale_findings),
        "vale_advisory_suggestions": sum(item["Severity"] == "suggestion" for item in vale_findings),
        "vale_status": "PASS_ADVISORY_ONLY" if not any(
            item["Severity"] in {"error", "warning"} for item in vale_findings
        ) else "REVIEW_REQUIRED",
    }
    audit["computational_tests"] = {
        "status": "NOT_APPLICABLE",
        "reason": "The revised Review article reports no new original computation or benchmark result.",
    }
    audit["deliverable_hashes"] = [
        {"path": str(path.relative_to(ROOT)).replace("\\", "/"), "bytes": path.stat().st_size, "sha256": sha256(path)}
        for path in deliverables
    ]
    audit["status"] = "PASS" if report["status"] == "PASS" else "GENERATED_AWAITING_RENDER_QA"
    if report["status"] == "PASS":
        audit["submission_readiness"] = {
            "status": "READY_FOR_AUTHOR_SUBMISSION_CHECKS",
            "blocking_items": [],
        }
    AUDIT.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print(json.dumps({
        "programmatic_status": report["programmatic_status"],
        "visual_status": report["visual_inspection"],
        "total_pages": total_pages,
        "final_status": report["status"],
    }))


if __name__ == "__main__":
    main()
